#!/usr/bin/env python3
"""Build the Round 1 one-slide submission on top of the official Teliport S4 template.

    pip install python-pptx
    python3 tools/build_round1_slide.py --case 3 \
        --template "1790163819038-Teliport S4_Case Study and Submission Template/Teliport S4_Submission Template.pptx" \
        --out submission/Round1_Case3_VeriQ.pptx

--case 3 builds VeriQ (adaptive QoS); --case 4 builds LabelTwin (medical labeling).

Edit TEAM / COLLEGE / MEMBERS below (or in PowerPoint afterwards).
"""
import argparse

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

TEAM = "<Team Name>"
COLLEGE = "<College>"
MEMBERS = "<Member 1> · <Member 2> · <Member 3>"

BLUE = RGBColor(0x0B, 0x5C, 0xAD)
NAVY = RGBColor(0x12, 0x23, 0x3D)
AMBER = RGBColor(0xE6, 0x8A, 0x00)
GREEN = RGBColor(0x1E, 0x8E, 0x5A)
INK = RGBColor(0x22, 0x2B, 0x38)
MUTED = RGBColor(0x55, 0x60, 0x70)
CARD = RGBColor(0xEE, 0xF5, 0xFC)
CARD2 = RGBColor(0xFF, 0xF6, 0xE6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def add_text(slide, x, y, w, h, paras, fill=None, line=None, anchor=MSO_ANCHOR.TOP,
             margin=0.08, shape=MSO_SHAPE.RECTANGLE):
    """paras: list of paragraphs; each paragraph is a list of (text, size, bold, color) runs,
    optionally prefixed with an alignment value."""
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = 0.06
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        align = PP_ALIGN.LEFT
        if para and not isinstance(para[0], tuple):
            align, para = para[0], para[1:]
        p.alignment = align
        p.space_after = Pt(1.5)
        for text, size, bold, color in para:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = "Calibri"
    return shp


def card(slide, x, y, w, h, title, body, fill=CARD, title_color=BLUE, size=10.5):
    paras = [[(title, 12.5, True, title_color)]]
    for item in body:
        if isinstance(item, list):
            paras.append(item)
        else:
            paras.append([("• ", size, True, title_color), (item, size, False, INK)])
    return add_text(slide, x, y, w, h, paras, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, margin=0.1)


def arrow(slide, x, y, w=0.22, h=0.32):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = AMBER
    a.line.fill.background()
    a.shadow.inherit = False


def build_cover(slide, case_no, tagline):
    for shp in slide.shapes:
        if not shp.has_text_frame:
            continue
        t = shp.text_frame.text.strip()
        if t.startswith("Case Study No"):
            shp.text_frame.paragraphs[0].runs[0].text = f"Case Study No. {case_no}"
            center = shp.left + shp.width // 2
            shp.width = Inches(5.5)
            shp.left = center - shp.width // 2
        elif t == "Team Details":
            shp.text_frame.paragraphs[0].runs[0].text = f"Team {TEAM}"
        elif shp.is_placeholder and shp.placeholder_format.type is not None and t == "":
            tf = shp.text_frame
            tf.text = f"{COLLEGE}\n{MEMBERS}\n{tagline}"
            for p in tf.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.size = Pt(18)
                    r.font.color.rgb = NAVY


def clear_placeholders(slide):
    # drop the empty layout placeholders; keep the template logos
    for shp in list(slide.shapes):
        if shp.is_placeholder and shp.has_text_frame:
            shp._element.getparent().remove(shp._element)


def header(slide, name, title, subtitle):
    add_text(slide, 2.62, 0.10, 9.55, 0.78, [
        [(name, 22, True, BLUE), (title, 17, True, NAVY)],
        [(subtitle + "   |   Team ", 10.5, False, MUTED),
         (TEAM, 10.5, True, MUTED), (" · " + COLLEGE, 10.5, False, MUTED)],
    ], anchor=MSO_ANCHOR.MIDDLE, margin=0.02)


def flow_strip(slide, y, h, label, boxes, dark=None):
    """A row of boxes joined by arrows; boxes are (width, title, body-or-lines)."""
    add_text(slide, 0.30, y - 0.02, 9.0, 0.28, [[(label, 11.5, True, NAVY)]], margin=0.0)
    x, gap, by, bh = 0.30, 0.22, y + 0.30, h - 0.32
    for i, (w, title, body) in enumerate(boxes):
        lines = body if isinstance(body, list) else [body]
        add_text(slide, x, by, w, bh, [[PP_ALIGN.CENTER, (title, 11, True, WHITE)]] +
                 [[PP_ALIGN.CENTER, (ln, 9, False, WHITE)] for ln in lines],
                 fill=NAVY if title == dark else BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                 anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
        x += w
        if i < len(boxes) - 1:
            arrow(slide, x + 0.01, by + bh / 2 - 0.15, w=gap - 0.03, h=0.30)
            x += gap


def footer(slide, paras):
    add_text(slide, 0.30, 6.80, 12.73, 0.58,
             [[(k, 9.5, True, NAVY), (v, 9.5, False, INK)] for k, v in paras],
             fill=RGBColor(0xF4, 0xF6, 0xF9), shape=MSO_SHAPE.ROUNDED_RECTANGLE,
             anchor=MSO_ANCHOR.MIDDLE, margin=0.08)


def build_case4(slide):
    clear_placeholders(slide)

    # header
    header(slide, "LabelTwin", ": Label-as-Code for self-validating, traceable medical device labels",
           "Case Study 4 · HLS · End-to-End Labeling Automation Management")

    # row A: problem / reframe / principles
    y, h = 0.98, 1.62
    card(slide, 0.30, y, 4.15, h, "The real problem", [
        "1 product change → 30+ labels × 20+ languages × packaging levels, re-verified by eye",
        "Values are copied across PLM, ERP, RIM & artwork, so errors are born at every hand-off",
        "Rules keep moving: ISO 20417:2026, ISO 15223-1 Amd 1:2025, FDA QMSR (Feb 2026), EUDAMED mandatory (2026)",
    ])
    card(slide, 4.58, y, 4.15, h, "Our reframe", [
        [("“Labeling is a data-consistency problem disguised as a document problem.”", 11, True, NAVY)],
        [("Every label becomes a computed, self-proving view of ", 10.5, False, INK),
         ("product data × market rules × language × packaging level", 10.5, True, INK),
         (". We make most errors structurally impossible, then let AI catch the rest.", 10.5, False, INK)],
    ], fill=CARD2, title_color=AMBER)
    card(slide, 8.86, y, 4.17, h, "Design principles", [
        "Deterministic core; GenAI & vision only at the unstructured edges",
        [("• ", 10.5, True, BLUE), ("AI may add findings, never clear them", 10.5, True, INK)],
        "Every element carries provenance + regulatory citation",
        "Humans decide at 3 e-signed gates (Part 11-style audit)",
        "Plugs into existing PLM / ERP / LMS: no rip-and-replace",
    ])

    # row B: architecture strip
    flow_strip(slide, 2.78, 1.62, "Architecture: agents propose, gates & humans dispose", [
        (1.80, "Sources", "PLM · ERP/SAP · RIM · Risk file & IFU · legacy artwork · FDA / EUR-Lex / MDCG / ISO feeds"),
        (2.05, "Knowledge layer", "Product golden record · Regulatory knowledge graph · Rule packs (Compliance-as-Code) · ISO 15223-1 symbol ontology"),
        (1.72, "Label Digital Twins", "Every element bound to source field + rule + hazard (ISO 14971)"),
        (2.78, "7 AI engines", ["① Reg Watcher   ② Change Impact (graph)", "③ GenAI Composer (RAG, cited)", "④ Compliance Checker (rules + LLM)", "⑤ Vision Verifier (OCR · symbols · GS1)", "⑥ Translation QA   ⑦ Risk Predictor"]),
        (1.74, "Governance", "Workflow state machine · e-signature · hash-chained audit · traceability matrix"),
        (1.62, "Outputs", "PRISYM / Loftware / Kallik · GUDID & EUDAMED · change order · line scan-to-verify"),
    ], dark="7 AI engines")

    # row C: innovations / proof
    yc, hc = 4.52, 2.18
    card(slide, 0.30, yc, 6.55, hc, "What makes it different", [
        [("1  Label Digital Twin: ", 10.5, True, NAVY),
         ("each element knows its data source, the clause requiring it (e.g. EU MDR Annex I §23.2(h) UDI) and the hazard it controls, so one click answers “Why is this here?”", 10.5, False, INK)],
        [("2  Compliance-as-Code: ", 10.5, True, NAVY),
         ("versioned, unit-tested rule packs (MDR §23, 21 CFR 801.18 / 801.20 / 801.109, GS1 UDI, ISO 15223-1 symbols) with effective dates", 10.5, False, INK)],
        [("3  Graph change-impact: ", 10.5, True, NAVY),
         ("a PLM or regulation change instantly yields impacted labels, markets, actions and a draft change order", 10.5, False, INK)],
        [("4  Vision + GS1 verification: ", 10.5, True, NAVY),
         ("auto-redline of artwork & print; DataMatrix AI(17)/(10) ↔ printed expiry/lot cross-check; translation invariants (37 °C never becomes 73 °C)", 10.5, False, INK)],
    ])
    card(slide, 6.98, yc, 6.05, hc, "How we will prove it: LabelBench", [
        "60 label twins · 5 markets (US, EU-DE, EU-FR, UK, CH) · 3 packaging levels; fictitious devices with GUDID-realistic attributes",
        "20+ defect types mined from FDA labeling-recall root causes, seeded with ground truth",
        "Ablation: rules-only vs LLM-only vs hybrid · timed user study vs manual checklist",
        "Trust built in: e-sign with meaning, hash-chained audit, risk-based tool validation (CSA)",
        [("Targets: ", 10.5, True, GREEN),
         ("≥95% recall on critical defects · change-impact days → minutes · ≥70% of repetitive checks automated · 100% element traceability", 10.5, True, INK)],
    ], fill=RGBColor(0xEC, 0xF7, 0xF1), title_color=GREEN)

    footer(slide, [
        ("Prototype stack: ", "FastAPI · Neo4j · YAML rule engine · PyMuPDF + PaddleOCR + OpenCV · zxing-cpp + GS1 parser · local LLM (Ollama) + bge-m3 RAG · React"),
        ("Fit for Tata Elxsi: ", "a tool-agnostic accelerator layered on the LMS/PLM stacks TE already implements for medtech clients, in line with the HLS “AI-first, regulatory-aware” direction. Data stays on-prem with a local LLM."),
    ])


def build_case3(slide):
    clear_placeholders(slide)
    header(slide, "VeriQ", ": intent-aware home QoS with measurable outcomes & explainable recovery",
           "Case Study 3 · MCV · Adaptive QoS Engine for Mixed Home Broadband Traffic")

    # row A: problem / idea / safety
    y, h = 0.98, 1.62
    card(slide, 0.30, y, 4.15, h, "The real problem", [
        "Queues, not bandwidth: a backup upload filled the modem's oversized FIFO and added 266 ms (p95) in our Day-0 lab (emulated, 5 Mbit up); fq_codel alone can't reach that queue",
        "QUIC/TLS on port 443 defeats port rules; apps rarely set DSCP",
        "Capacity moves (DOCSIS, 5G FWA, Wi-Fi): a static shaper is bloated or wasteful",
    ])
    card(slide, 4.58, y, 4.15, h, "Our idea", [
        [("“VeriQ finds the cause, protects what you chose, and proves it.”", 11, True, NAVY)],
        [("Localise the fault first", 10.5, True, INK),
         (" (own queue, WAN drop, remote path, Wi-Fi, bad probe), then act with bounded fast guardrails. Every adaptation gets an ", 10.5, False, INK),
         ("experience receipt: improved / degraded / inconclusive", 10.5, True, INK),
         (", measured on real calls and video.", 10.5, False, INK)],
    ], fill=CARD2, title_color=AMBER)
    card(slide, 8.86, y, 4.17, h, "Bounded by design", [
        "Never inspects forwarded payload; headers + timing; IPv4 + IPv6",
        [("• ", 10.5, True, BLUE), ("Intents = expiring, budgeted leases; CAKE caps any boost; misbehaving flows sanctioned; progress measured", 10.5, True, INK)],
        "Abstains when unsure: never promotes an uncertain flow",
        "Shadow mode first; every decision logged; watchdog → SAFE shaper",
    ])

    # row B: closed loop
    flow_strip(slide, 2.78, 1.62, "Three speeds: fast guardrails (sub-second) · outcome receipts (tens of seconds) · controlled experiments (lab, for claims)", [
        (1.86, "Observe", "Header + timing of first packets per flow, app marks (NQB / L4S); v4/v6; no payload"),
        (2.12, "Understand", "Marks first, then a calibrated classifier (abstains) · capacity μ ± σ + UDPST idle test (RFC 9946) · fault localiser"),
        (2.04, "Decide", "Bounded guardrails: shaper = μ − kσ, change-point fast path · leases within budget"),
        (1.86, "Enforce", "CAKE with bounded lease boosts; ordered, versioned, read-back apply"),
        (2.04, "Assess", "Receipts on WebRTC freezes, DASH rebuffers, delay, progress; estimated MOS labelled"),
        (1.71, "Show", "Grafana · 3D twin + replay: measured / estimated / illustrative"),
    ], dark="Assess")

    # row C: innovations / proof
    yc, hc = 4.52, 2.18
    card(slide, 0.30, yc, 6.55, hc, "What makes it different", [
        [("1  Fault localisation before action: ", 10.5, True, NAVY),
         ("WAN drop vs remote path vs Wi-Fi vs lying probe, scored against emulator ground truth; a remote slowdown never triggers a wasteful cut", 10.5, False, INK)],
        [("2  Real-app evidence: ", 10.5, True, NAVY),
         ("real encrypted WebRTC calls and DASH players in headless Chromium; we count video freezes and rebuffers, not just pings", 10.5, False, INK)],
        [("3  Payload-free, calibrated classification: ", 10.5, True, NAVY),
         ("marks first; split by flow and by time, no shortcut features (SIGCOMM 2025 lesson); abstains when unsure", 10.5, False, INK)],
        [("4  Honest receipts + replay: ", 10.5, True, NAVY),
         ("improved / degraded / inconclusive, plus IETF Quality-of-Outcome scores; 3D replay shows inferred vs true ISP queue", 10.5, False, INK)],
    ])
    card(slide, 6.98, yc, 6.05, hc, "How we will prove it", [
        "Real, tuned baselines: FIFO · fq_codel · SQM-CAKE · cake-autorate · qosify-style rules + ablations; ≥10 randomised runs, 95% CIs",
        "13 scenarios: ISO during call, 100 → 20 Mbps drop, 3 TVs + gaming, remote delay, Wi-Fi, false marks, crash, IPv6…",
        "Physical Pi 5 kit: the judge picks the fault, VeriQ responds from telemetry alone",
        [("Targets: ", 10.5, True, GREEN),
         ("0 call freezes in S1 · added delay p95 ≤ 10 ms · settle < 2 s after WAN drop · measured progress for every class", 10.5, True, INK)],
    ], fill=RGBColor(0xEC, 0xF7, 0xF1), title_color=GREEN)

    footer(slide, [
        ("Stack: ", "Linux tc (CAKE; HTB later) · Python + pyroute2 control plane · LightGBM + conformal · Prometheus/Grafana · React + three.js · Playwright (WebRTC, dash.js) · iperf3 / irtt / Flent"),
        ("Industry path (designed, not built): ", "RDK-B / OpenWrt package · TR-181 + TR-369 USP fleet management · research tracks: eBPF datapath, Wi-Fi airtime, L4S, multi-gig scaling"),
    ])


CASES = {
    3: (build_case3, "VeriQ: adapts, measures, explains"),
    4: (build_case4, "LabelTwin: Label-as-Code for medical device labeling"),
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--case", type=int, choices=sorted(CASES), default=3)
    ap.add_argument("--template", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    prs = Presentation(a.template)
    builder, tagline = CASES[a.case]
    build_cover(prs.slides[0], a.case, tagline)
    builder(prs.slides[1])
    prs.save(a.out)
    print("wrote", a.out)


if __name__ == "__main__":
    main()
